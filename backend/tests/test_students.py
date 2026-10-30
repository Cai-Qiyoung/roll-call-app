"""名单模块单元测试：Excel/CSV 导入、模板下载、删除与清空。"""

from io import BytesIO

from openpyxl import Workbook, load_workbook

from models import RollRecord, Student

from helpers import add_student


def _xlsx_bytes(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["学号", "姓名", "专业"])
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_xlsx_creates_students(client, db):
    content = _xlsx_bytes(
        [["001", "张三", "软件工程"], ["002", "李四", "计算机"], ["003", "王五", "人工智能"]]
    )
    res = client.post(
        "/students/import",
        files={"file": ("list.xlsx", content, "application/octet-stream")},
    )
    assert res.json()["imported"] == 3
    assert db.query(Student).count() == 3


def test_import_csv(client, db):
    content = "学号,姓名,专业\n001,张三,软件工程\n002,李四,计算机\n".encode("utf-8-sig")
    res = client.post(
        "/students/import",
        files={"file": ("list.csv", content, "text/csv")},
    )
    assert res.json()["imported"] == 2


def test_reimport_updates_existing_student(client, db):
    add_student(db, "001", name="旧名字")
    content = _xlsx_bytes([["001", "新名字", "新专业"]])
    client.post(
        "/students/import",
        files={"file": ("list.xlsx", content, "application/octet-stream")},
    )
    db.refresh(db.query(Student).filter(Student.student_no == "001").first())
    assert db.query(Student).one().name == "新名字"


def test_template_download_is_valid_xlsx(client, db):
    res = client.get("/students/template")
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    assert wb.active.cell(1, 1).value == "学号"


def test_delete_student_removes_records(client, db):
    s = add_student(db, "001")
    db.add(RollRecord(student_id=s.id, kind="attend", delta=1))
    db.commit()

    res = client.delete("/students/001")
    assert res.status_code == 200
    assert db.query(Student).count() == 0
    assert db.query(RollRecord).count() == 0


def test_clear_all_resets_everything(client, db):
    add_student(db, "001")
    add_student(db, "002")
    client.post("/rollcall/sequence")  # 游标走到第 1 位

    res = client.delete("/students")
    assert res.json()["cleared"] is True
    assert db.query(Student).count() == 0

    # 清空后重新导入，顺序点名从第一位开始
    client.post(
        "/students/import",
        files={"file": ("l.csv", "学号,姓名,专业\nA1,A1x,软工".encode("utf-8-sig"), "text/csv")},
    )
    assert client.post("/rollcall/sequence").json()["student_no"] == "A1"