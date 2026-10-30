"""排名模块单元测试：Top 排序、统计、详单导出。"""

from io import BytesIO

from openpyxl import load_workbook

from helpers import add_student


def test_top_orders_by_score_desc(client, db):
    add_student(db, "001", score=1)
    add_student(db, "002", score=3)
    add_student(db, "003", score=2)
    top = client.get("/ranking/top", params={"n": 10}).json()
    assert [t["student_no"] for t in top] == ["002", "003", "001"]
    assert top[0]["rank"] == 1


def test_top_limits_count(client, db):
    for i in range(15):
        add_student(db, f"0{i}", score=i)
    assert len(client.get("/ranking/top", params={"n": 5}).json()) == 5


def test_stats(client, db):
    add_student(db, "001", score=2)
    add_student(db, "002", score=4)
    stats = client.get("/ranking/stats").json()
    assert stats["total_students"] == 2
    assert stats["max_score"] == 4
    assert stats["avg_score"] == 3.0


def test_export_detail_xlsx_header_and_rows(client, db):
    add_student(db, "001", name="张三", score=5, roll_count=3)
    add_student(db, "002", name="李四", score=2, roll_count=1)

    res = client.get("/ranking/export")
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    ws = wb.active
    assert [ws.cell(1, c).value for c in range(1, 7)] == [
        "排名", "学号", "姓名", "专业", "随机点名次数", "总积分",
    ]
    assert ws.max_row == 3  # 表头 + 2 行数据
    assert ws.cell(2, 2).value == "001"  # 积分高者排第一