import csv
from io import BytesIO, StringIO

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from database import get_db
from models import AppState, RollRecord, Student
from schemas import StudentOut

router = APIRouter(prefix="/students", tags=["students"])


@router.get("", response_model=list[StudentOut])
def list_students(db: Session = Depends(get_db)):
    return db.query(Student).order_by(Student.student_no).all()


@router.post("/import")
async def import_students(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """从 Excel(.xlsx) 或 CSV 导入学生名单，表头顺序：学号、姓名、专业。"""
    content = await file.read()
    rows = []  # [(学号, 姓名, 专业), ...]
    filename = (file.filename or "").lower()

    if filename.endswith(".csv"):
        text = content.decode("utf-8-sig")
        for r in csv.DictReader(StringIO(text)):
            rows.append(
                (
                    str(r.get("学号", "") or "").strip(),
                    str(r.get("姓名", "") or "").strip(),
                    str(r.get("专业", "") or "").strip(),
                )
            )
    else:
        wb = load_workbook(BytesIO(content))
        ws = wb.active
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or r[0] is None:
                continue
            no = str(r[0]).strip()
            nm = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ""
            major = str(r[2]).strip() if len(r) > 2 and r[2] is not None else ""
            rows.append((no, nm, major))

    count = 0
    for no, nm, major in rows:
        if not no or not nm:
            continue
        student = db.query(Student).filter(Student.student_no == no).first()
        if student:
            student.name = nm
            student.major = major
        else:
            db.add(Student(student_no=no, name=nm, major=major))
        count += 1
    db.commit()
    return {"imported": count}


@router.get("/template")
def download_template():
    """生成导入模板（学号、姓名、专业 + 两行示例）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "学生名单"
    ws.append(["学号", "姓名", "专业"])
    ws.append(["2024001", "张三", "软件工程"])
    ws.append(["2024002", "李四", "软件工程"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=student_template.xlsx"},
    )


@router.delete("/{student_no}")
def delete_student(student_no: str, db: Session = Depends(get_db)):
    student = db.query(Student).filter(Student.student_no == student_no).first()
    if not student:
        raise HTTPException(status_code=404, detail="学生不存在")
    db.query(RollRecord).filter(RollRecord.student_id == student.id).delete()
    db.delete(student)
    db.commit()
    return {"deleted": student_no}


@router.delete("")
def clear_students(db: Session = Depends(get_db)):
    """清空名单与全部记录，并将顺序点名游标归零。"""
    db.query(RollRecord).delete()
    db.query(Student).delete()
    state = db.get(AppState, "seq_index")
    if state:
        state.value = "0"
    db.commit()
    return {"cleared": True}